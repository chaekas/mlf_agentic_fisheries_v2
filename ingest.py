"""
MLF Fisheries Agentic AI - Universal Document Ingestion

Supported:
PDF, DOCX, XLSX/XLSM, CSV, PPTX, TXT, MD, JSON, XML, HTML and common images.

Embeddings:
BAAI/bge-m3

Vector database:
PostgreSQL/pgvector (including Supabase PostgreSQL)
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
from typing import Iterable

import pandas as pd
from dotenv import load_dotenv
from langchain_core.documents import Document
from langchain_postgres import PGVector
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PIL import Image
from pptx import Presentation
from PyPDF2 import PdfReader
from docx import Document as DocxDocument

try:
    import pytesseract
except ImportError:
    pytesseract = None

# Avoid accidental transformer image imports where supported.
os.environ.setdefault("TRANSFORMERS_NO_TORCHVISION", "1")

from langchain_huggingface import HuggingFaceEmbeddings


load_dotenv()

DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
COLLECTION = os.getenv("PGVECTOR_COLLECTION", "mlf_fisheries_regulations")
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "BAAI/bge-m3")

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".pptx",
    ".txt", ".md", ".json", ".xml", ".html", ".htm",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif",
}

splitter = RecursiveCharacterTextSplitter(
    chunk_size=900,
    chunk_overlap=120,
    separators=["\n\n", "\n", ". ", " ", ""],
)


def validate_config() -> None:
    if not DATABASE_URL:
        raise RuntimeError(
            "DATABASE_URL is missing. Set it in .env, for example:\n"
            "DATABASE_URL=postgresql://USER:PASSWORD@HOST:5432/postgres"
        )


def detect_lang(text: str) -> str:
    """Detect Swahili vs English without failing ingestion if detection is uncertain."""
    try:
        from langdetect import detect
        result = detect(text[:1500])
        return "sw" if result.startswith("sw") else "en"
    except Exception:
        return "en"


def file_id(path: Path) -> str:
    """Stable ID used to identify a source file."""
    return hashlib.sha256(str(path.resolve()).encode("utf-8")).hexdigest()[:16]


def text_document(
    text: str,
    source: str,
    metadata: dict | None = None,
) -> list[Document]:
    text = (text or "").strip()
    if not text:
        return []

    base = {
        "source": source,
        "lang": detect_lang(text),
    }
    if metadata:
        base.update(metadata)

    docs: list[Document] = []
    for chunk_no, chunk in enumerate(splitter.split_text(text), 1):
        chunk = chunk.strip()
        if not chunk:
            continue
        item = dict(base)
        item["chunk"] = chunk_no
        docs.append(Document(page_content=chunk, metadata=item))
    return docs


def load_pdf(path: Path) -> list[Document]:
    docs: list[Document] = []
    reader = PdfReader(str(path))
    fid = file_id(path)

    for page_no, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        docs.extend(
            text_document(
                text,
                path.name,
                {
                    "file_id": fid,
                    "document": path.name,
                    "file_type": "pdf",
                    "page": page_no,
                },
            )
        )
    return docs


def load_docx(path: Path) -> list[Document]:
    document = DocxDocument(str(path))
    paragraphs = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip()
    ]

    table_rows: list[str] = []
    for table_no, table in enumerate(document.tables, 1):
        table_rows.append(f"Table {table_no}:")
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                table_rows.append(row_text)

    text = "\n".join(paragraphs + table_rows)
    return text_document(
        text,
        path.name,
        {
            "file_id": file_id(path),
            "document": path.name,
            "file_type": "docx",
        },
    )


def load_spreadsheet(path: Path) -> list[Document]:
    sheets = pd.read_excel(path, sheet_name=None)
    docs: list[Document] = []
    fid = file_id(path)

    for sheet_name, frame in sheets.items():
        frame = frame.fillna("")
        text = f"Sheet: {sheet_name}\n{frame.to_csv(index=False)}"
        docs.extend(
            text_document(
                text,
                path.name,
                {
                    "file_id": fid,
                    "document": path.name,
                    "file_type": "spreadsheet",
                    "sheet": str(sheet_name),
                },
            )
        )
    return docs


def load_csv(path: Path) -> list[Document]:
    frame = pd.read_csv(path).fillna("")
    return text_document(
        frame.to_csv(index=False),
        path.name,
        {
            "file_id": file_id(path),
            "document": path.name,
            "file_type": "csv",
        },
    )


def load_pptx(path: Path) -> list[Document]:
    presentation = Presentation(str(path))
    docs: list[Document] = []
    fid = file_id(path)

    for slide_no, slide in enumerate(presentation.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())

        docs.extend(
            text_document(
                "\n".join(texts),
                path.name,
                {
                    "file_id": fid,
                    "document": path.name,
                    "file_type": "pptx",
                    "slide": slide_no,
                },
            )
        )
    return docs


def load_plain(path: Path) -> list[Document]:
    return text_document(
        path.read_text(encoding="utf-8", errors="ignore"),
        path.name,
        {
            "file_id": file_id(path),
            "document": path.name,
            "file_type": path.suffix.lower().lstrip("."),
        },
    )


def load_json(path: Path) -> list[Document]:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return text_document(
        json.dumps(data, ensure_ascii=False, indent=2),
        path.name,
        {
            "file_id": file_id(path),
            "document": path.name,
            "file_type": "json",
        },
    )


def load_image(path: Path) -> list[Document]:
    if pytesseract is None:
        print(
            f"WARNING: pytesseract is not installed; OCR skipped for {path.name}."
        )
        return []

    try:
        text = pytesseract.image_to_string(Image.open(path))
        return text_document(
            text,
            path.name,
            {
                "file_id": file_id(path),
                "document": path.name,
                "file_type": "image",
                "ocr": True,
            },
        )
    except Exception as exc:
        print(f"WARNING: OCR failed for {path.name}: {exc}")
        return []


def load_file(path: Path) -> list[Document]:
    ext = path.suffix.lower()

    loaders = {
        ".pdf": load_pdf,
        ".docx": load_docx,
        ".xlsx": load_spreadsheet,
        ".xlsm": load_spreadsheet,
        ".csv": load_csv,
        ".pptx": load_pptx,
        ".txt": load_plain,
        ".md": load_plain,
        ".xml": load_plain,
        ".html": load_plain,
        ".htm": load_plain,
        ".json": load_json,
        ".png": load_image,
        ".jpg": load_image,
        ".jpeg": load_image,
        ".webp": load_image,
        ".bmp": load_image,
        ".tiff": load_image,
        ".tif": load_image,
    }

    loader = loaders.get(ext)
    return loader(path) if loader else []


def collect_paths(items: Iterable[str]) -> list[Path]:
    paths: list[Path] = []

    for item in items:
        path = Path(item)
        if path.is_dir():
            paths.extend(
                p for p in path.rglob("*")
                if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
            )
        elif path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            paths.append(path)
        else:
            print(f"Skipping unsupported or missing path: {path}")

    # Remove duplicates while preserving order.
    seen = set()
    result = []
    for p in paths:
        key = str(p.resolve()).lower()
        if key not in seen:
            seen.add(key)
            result.append(p)
    return result


def get_embeddings() -> HuggingFaceEmbeddings:
    try:
        return HuggingFaceEmbeddings(
            model_name=EMBEDDING_MODEL,
            encode_kwargs={"normalize_embeddings": True},
            model_kwargs={"device": os.getenv("EMBEDDING_DEVICE", "cpu")},
        )
    except Exception as exc:
        raise RuntimeError(
            "Could not load the BGE-M3 embedding model. "
            "If the error mentions torchvision, install a torchvision version "
            "compatible with your installed PyTorch, then restart the virtual environment."
        ) from exc


def get_vectorstore() -> PGVector:
    validate_config()
    return PGVector(
        embeddings=get_embeddings(),
        collection_name=COLLECTION,
        connection=DATABASE_URL,
        use_jsonb=True,
        create_extension=False,
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Index MLF Fisheries documents into PostgreSQL/pgvector."
    )
    parser.add_argument(
        "paths",
        nargs="+",
        help="Files or directories to ingest",
    )
    args = parser.parse_args()

    paths = collect_paths(args.paths)
    if not paths:
        print("No supported files found.")
        return

    documents: list[Document] = []

    for path in paths:
        print(f"Reading: {path}")
        try:
            loaded = load_file(path)
            documents.extend(loaded)
            print(f"  -> {len(loaded)} chunks")
        except Exception as exc:
            print(f"  ERROR: {exc}")

    if not documents:
        print("No text could be extracted from the supplied files.")
        return

    print(f"Creating embeddings with {EMBEDDING_MODEL}...")
    store = get_vectorstore()

    print(f"Adding {len(documents)} chunks to collection '{COLLECTION}'...")
    store.add_documents(documents)

    print("\nSUCCESS")
    print(f"Indexed chunks : {len(documents)}")
    print(f"Collection     : {COLLECTION}")
    print(f"Database host  : configured through DATABASE_URL")


if __name__ == "__main__":
    main()
